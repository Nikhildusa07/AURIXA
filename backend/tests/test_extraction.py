from pathlib import Path

import pytest

from app.documents.extraction import extract_text


def test_extract_txt_file(tmp_path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text(
        "Hello AURIXA document processing",
        encoding="utf-8",
    )

    result = extract_text(file_path)

    assert result == "Hello AURIXA document processing"


def test_file_not_found():
    with pytest.raises(FileNotFoundError):
        extract_text("non_existent_document.txt")


def test_path_is_not_a_file(tmp_path):
    with pytest.raises(
        ValueError,
        match="Path is not a file",
    ):
        extract_text(tmp_path)


def test_unsupported_file_type(tmp_path):
    file_path = tmp_path / "sample.csv"
    file_path.write_text("test", encoding="utf-8")

    with pytest.raises(
        ValueError,
        match="Unsupported document type",
    ):
        extract_text(file_path)


def test_extract_docx_file(tmp_path):
    from docx import Document

    file_path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(file_path)

    result = extract_text(file_path)

    assert "First paragraph" in result
    assert "Second paragraph" in result